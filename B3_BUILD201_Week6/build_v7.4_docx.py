#!/usr/bin/env python3
"""v7.4 — student LMS docx + separate writing-center handout docx."""
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parent
STUDENT_OUT = ROOT / "BUILD201_Week6_Student_v7.4.docx"
HANDOUT_OUT = ROOT / "BUILD201_Writing_Center_Handout_v7.4.docx"

BRIDGE = (
    "When I revise with AI‑assisted drafting（especially for clarity）, "
    "I stop here：I check whether the file still states what I intend."
)
SAMPLE = (
    "Strong recall. Next time, name the trade-off earlier. "
    + BRIDGE
    + " Tie the limitation to workflow, not only opinion."
)


def student_doc():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("BUILD 201 — Week 6 Reflection", level=0)
    for line in (ROOT / "ASSIGNMENT_BRIEF_v7.4_student.md").read_text().splitlines():
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.strip():
            doc.add_paragraph(line.strip().replace("**", ""))
    doc.add_page_break()
    doc.add_heading("Feedback rubric", level=1)
    for line in (ROOT / "FEEDBACK_RUBRIC_v7.4_student.md").read_text().splitlines():
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("|") or line.startswith("---"):
            continue
        elif line.strip() and not line.startswith(">"):
            doc.add_paragraph(line.strip().replace("**", ""))
    doc.save(STUDENT_OUT)


def handout_doc():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph(
        "For writing center staff — not distributed on the student assignment page."
    ).runs[0].italic = True
    for line in (ROOT / "WRITING_CENTER_HANDOUT_v7.4.md").read_text().splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.add_run(line[2:]).italic = True
        elif line.strip():
            doc.add_paragraph(line.strip().replace("**", ""))
    doc.save(HANDOUT_OUT)


def main():
    student_doc()
    handout_doc()
    print(f"wrote {STUDENT_OUT} ({STUDENT_OUT.stat().st_size} bytes)")
    print(f"wrote {HANDOUT_OUT} ({HANDOUT_OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
